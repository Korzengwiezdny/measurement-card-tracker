import base64, http.cookiejar, io, json, sqlite3, secrets, threading, unittest, urllib.request, urllib.error
from pathlib import Path
import openpyxl
import server

class AppTests(unittest.TestCase):
    def setUp(self):
        folder=Path(__file__).resolve().parent/'work';folder.mkdir(exist_ok=True)
        self.path=folder/('test-'+secrets.token_hex(8)+'.sqlite3');server.DB_PATH=self.path;server.init_db();server.SESSIONS.clear();server.ATTEMPTS.clear()
        self.user={'name':'technik'}
        base=dict(mu='MU1',serial='SN1',manufacturer='NI',model='PXIe 4081',hall='Maintenance',location='Magazyn',designation='D',calibration='2030-01-01',calibration_note='',replacement='',notes='',source_row=4,review='',version=0,status='sprawna')
        with server.connect() as c:server.insert_cards(c,[dict(base,id='a'),dict(base,id='b',mu='MU2',serial='SN2',hall='Test C4',location='ATE 08'),dict(base,id='c',mu='MU3',serial='SN3')])
        self.b=dict(incoming='a',outgoing='b',incomingVersion=0,outgoingVersion=0,destinationHall='Maintenance',destinationLocation='Kalibracja',note='Test',incomingVerified=True,outgoingVerified=True,destinationVerified=True,destinationStatus='kalibracja',calibrationVerified=True)
    def tearDown(self):
        for suffix in ('','-wal','-shm'):
            Path(str(self.path)+suffix).unlink(missing_ok=True)
    def test_new_register_archives_and_replaces(self):
        with server.connect() as c:
            c.execute("INSERT INTO users(name,password,role) VALUES('admin','hash','admin')");c.commit()
            row=dict(c.execute("SELECT * FROM cards WHERE id='a'").fetchone());row.update(id='new',serial='NEW-SN')
            revision=server.register_revision(c)
            result=server.replace_register(c,[row],dict(confirmNewRegister=True,registerName='Nowy',revision=revision,filename='test.xlsx'),self.user)
            self.assertEqual(c.execute('SELECT id FROM cards').fetchall()[0][0],'new')
            self.assertEqual(c.execute('SELECT COUNT(*) FROM users').fetchone()[0],1)
            self.assertEqual(c.execute('SELECT kind FROM events').fetchone()[0],'nowy spis')
            archive=server.DB_PATH.parent.parent/'backups'/result['backup']
            try:
                with sqlite3.connect(archive,factory=server.Connection) as old:self.assertEqual(old.execute('SELECT COUNT(*) FROM cards').fetchone()[0],3)
            finally:archive.unlink()
    def test_new_register_rejects_stale_preview(self):
        with server.connect() as c:
            revision=server.register_revision(c)
            c.execute("UPDATE cards SET version=1 WHERE id='a'");c.commit()
            with self.assertRaises(server.Problem):server.replace_register(c,[],dict(confirmNewRegister=True,registerName='Nowy',revision=revision),self.user)
            self.assertEqual(c.execute('SELECT COUNT(*) FROM cards').fetchone()[0],3)
    def test_new_register_requires_confirmation(self):
        with server.connect() as c:
            with self.assertRaises(server.Problem):server.replace_register(c,[],dict(registerName='Nowy'),self.user)
            self.assertEqual(c.execute('SELECT COUNT(*) FROM cards').fetchone()[0],3)

    def test_calibration_matching_and_excel(self):
        with server.connect() as c:
            c.execute("UPDATE cards SET serial='PART1 + PART2',mu='' WHERE id='a'");c.commit()
            self.assertEqual([r['id'] for r in server.calibration_matches(c,'',' part2 ')],['a'])
            self.assertEqual(server.calibration_matches(c,'MU2','PART1'),[])
            wb=openpyxl.Workbook();ws=wb.active;ws.append(['SN','MU']);ws.append(['PART2','']);raw=io.BytesIO();wb.save(raw)
            result=server.calibration_preview(c,{'file':base64.b64encode(raw.getvalue()).decode()})
            self.assertEqual(result['rows'][0]['matches'][0]['id'],'a')
    def test_calibration_swap_undo_and_duplicate(self):
        with server.connect() as c:
            body=dict(batch='Kalibracja',ids=['b'])
            self.assertEqual(server.calibration_add(c,body,self.user)['added'],1)
            self.assertEqual(server.calibration_add(c,body,self.user)['skipped'],1)
            event=server.swap(c,{**self.b,'destinationStatus':'kalibracja'},self.user)
            plan=server.calibration_list(c)[0];self.assertTrue(plan['swapAt']);self.assertEqual(plan['card']['location'],'Kalibracja')
            server.undo(c,{'id':event,'note':'Test'},self.user)
            self.assertEqual(server.calibration_list(c)[0]['swapAt'],'')
            server.calibration_close(c,dict(id=plan['id'],version=0,state='anulowane',reason='Odwołane'),self.user)
            self.assertEqual(server.calibration_add(c,body,self.user)['added'],1)
    def test_calibration_ambiguous_serial_not_auto_resolved(self):
        with server.connect() as c:
            c.execute("UPDATE cards SET serial='SN1' WHERE id='c'");c.commit()
            self.assertEqual(len(server.calibration_matches(c,'','SN1')),2)
            self.assertEqual(server.calibration_matches(c,'','UNKNOWN'),[])

    def new_card_body(self,**patch):
        return dict(serial='NEW1 + NEW2',model='PXIe 4081',hall='Maintenance',location='Magazyn MD',status='nieustalony',physicallyVerified=True,reason='Zakup 2026',**patch)
    def test_stock_shared_spare_and_model_alias(self):
        with server.connect() as c:
            c.execute("UPDATE locations SET stock=1 WHERE name='Magazyn'")
            c.execute("UPDATE cards SET model='PXI-E 4081' WHERE id='a'")
            c.execute("UPDATE cards SET location='ATE 09' WHERE id='c'");c.commit()
            report=server.stock_report(c)
            self.assertEqual(report['ready'],1)
            self.assertEqual(len(report['stations']),2)
            self.assertTrue(all(not s['missing'] for s in report['stations']))
    def test_stock_excludes_broken_duplicate_and_service(self):
        with server.connect() as c:
            c.execute("UPDATE locations SET stock=1 WHERE name='Magazyn'")
            c.execute("UPDATE cards SET status='naprawa' WHERE id='a'")
            c.execute("UPDATE cards SET serial='SN2' WHERE id='c'");c.commit()
            report=server.stock_report(c)
            self.assertEqual(report['ready'],0);self.assertEqual(len(report['missingModels']),1)
            self.assertEqual(len(report['models'][0]['blocked']),2)
    def test_stock_multiple_serials_are_one_spare(self):
        with server.connect() as c:
            c.execute("UPDATE locations SET stock=1 WHERE name='Magazyn'")
            c.execute("UPDATE cards SET serial='PART1 + PART2',mu='' WHERE id='a'")
            c.execute("UPDATE cards SET retired=1 WHERE id='c'");c.commit()
            report=server.stock_report(c);self.assertEqual((report['stored'],report['ready']),(1,1))
    def test_add_empty_station_then_assign_card(self):
        with server.connect() as c:
            station=server.add_station(c,dict(name=' ate16 ',hall='H2'),self.user)
            self.assertEqual(station['name'],'ATE 16')
            self.assertEqual(c.execute('SELECT count(*) FROM cards').fetchone()[0],3)
            body={**self.new_card_body(),'location':'ATE 16','hall':'H2','locationId':station['id'],'status':'sprawna','calibration':'2030-01-01','installationVerified':True}
            card=server.register_card(c,body,self.user)
            self.assertEqual(c.execute('SELECT location,hall FROM cards WHERE id=?',(card['id'],)).fetchone()[:],('ATE 16','H2'))
    def test_station_duplicate_alias_is_blocked(self):
        with server.connect() as c:server.add_station(c,dict(name='ATE 1',hall='H2'),self.user)
        for name in ('ate01','ATE 01','ATE 1','ATE0001'):
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.add_station(c,dict(name=name,hall='H3'),self.user)
    def test_station_invalid_name_has_no_side_effect(self):
        for name in ('Magazyn','ATE','ATE -1','ATE test'):
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.add_station(c,dict(name=name,hall='H2'),self.user)
        with server.connect() as c:self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],0)
    def move_body(self,**patch):
        with server.connect() as c:lid=c.execute("SELECT id FROM locations WHERE name='Dział kalibracji'").fetchone()[0]
        return {**dict(id='a',version=0,locationId=lid,status='kalibracja',reason='Oddanie do kalibracji',physicallyVerified=True),**patch}
    def test_move_and_service_return_are_audited(self):
        with server.connect() as c:
            server.move_card(c,self.move_body(),self.user)
            row=dict(c.execute("SELECT * FROM cards WHERE id='a'").fetchone())
            self.assertEqual((row['hall'],row['location'],row['status'],row['version']),('H2','Dział kalibracji','kalibracja',1))
            lid=c.execute("SELECT id FROM locations WHERE name='Magazyn MD'").fetchone()[0]
            server.move_card(c,self.move_body(version=1,locationId=lid,operation='powrót z kalibracji',result='pozytywny',document='CERT/123',calibration='2031-01-01',status='sprawna'),self.user)
            row=dict(c.execute("SELECT * FROM cards WHERE id='a'").fetchone())
            self.assertEqual(row['calibration'],'2031-01-01');self.assertIn('CERT/123',row['calibration_note'])
            self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],2)
    def test_move_invalid_requests_do_not_change_cards(self):
        for patch in [dict(version=99),dict(physicallyVerified=False),dict(locationId=-1),dict(status='wycofana'),dict(operation='powrót z naprawy',result='negatywny',document='SER/1',status='sprawna')]:
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.move_card(c,self.move_body(**patch),self.user)
        with server.connect() as c:self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],0)
    def test_location_disable_blocks_future_moves_but_preserves_card(self):
        body=self.move_body()
        with server.connect() as c:
            server.move_card(c,body,self.user)
            server.location_update(c,dict(id=body['locationId'],name='Dział kalibracji',hall='H2',active=False),self.user)
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='a'").fetchone()[0],'Dział kalibracji')
        with self.assertRaises(server.Problem):
            with server.connect() as c:server.move_card(c,{**body,'id':'b'},self.user)
    def test_negative_service_result_keeps_card_active(self):
        with server.connect() as c:
            server.move_card(c,self.move_body(operation='powrót z naprawy',result='negatywny',document='SER/FAIL',status='do sprawdzenia'),self.user)
            row=dict(c.execute("SELECT * FROM cards WHERE id='a'").fetchone())
            self.assertEqual((row['retired'],row['status']),(0,'do sprawdzenia'))
    def test_new_installation_requires_calibration_and_confirmation(self):
        with server.connect() as c:
            c.execute("INSERT INTO locations(name,hall) VALUES('ATE 99','Test')");lid=c.execute('SELECT last_insert_rowid()').fetchone()[0];c.commit()
        body=self.move_body(locationId=lid,status='sprawna')
        for patch in [{},dict(calibrationVerified=True),dict(installationVerified=True)]:
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.move_card(c,{**body,**patch},self.user)
        with server.connect() as c:server.move_card(c,{**body,'calibrationVerified':True,'installationVerified':True},self.user)
    def test_add_without_mu_and_audit(self):
        with server.connect() as c:
            result=server.register_card(c,self.new_card_body(),self.user)
            row=dict(c.execute('SELECT * FROM cards WHERE id=?',(result['id'],)).fetchone())
            self.assertEqual((row['mu'],row['location'],row['retired']),('','Magazyn MD',0))
            event=c.execute('SELECT * FROM events').fetchone()
            self.assertEqual(json.loads(event['before_json']),[])
            self.assertEqual(json.loads(event['after_json'])[0]['serial'],'NEW1 + NEW2')
    def test_add_duplicate_component_and_missing_confirmation_rejected(self):
        for patch in [dict(serial='NEW + SN1'),dict(physicallyVerified=False),dict(model='')]:
            body={**self.new_card_body(),**patch}
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.register_card(c,body,self.user)
        with server.connect() as c:self.assertEqual(c.execute('SELECT count(*) FROM cards').fetchone()[0],3)
    def test_add_directly_to_ate_checks_installation(self):
        body={**self.new_card_body(),'location':'ATE 08','status':'sprawna','calibration':'2030-01-01','installationVerified':True}
        for patch in [dict(status='nieustalony'),dict(calibration='2020-01-01'),dict(calibration=''),dict(installationVerified=False)]:
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.register_card(c,{**body,**patch},self.user)
        with server.connect() as c:
            result=server.register_card(c,body,self.user)
            self.assertEqual(c.execute('SELECT location FROM cards WHERE id=?',(result['id'],)).fetchone()[0],'ATE 08')
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='b'").fetchone()[0],'ATE 08')
    def test_retirement_preserves_history_and_blocks_use(self):
        body=dict(id='a',version=0,hall='Maintenance',location='Zlom / izolacja',reason='Brak mozliwosci naprawy',physicallyVerified=True)
        for patch in [dict(reason=''),dict(version=9),dict(physicallyVerified=False),dict(location='ATE 08')]:
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.retire_card(c,{**body,**patch},self.user)
        with server.connect() as c:
            server.retire_card(c,body,self.user)
            row=dict(c.execute("SELECT * FROM cards WHERE id='a'").fetchone())
            self.assertEqual((row['retired'],row['status'],row['version']),(1,'wycofana',1))
            self.assertEqual(c.execute('SELECT count(*) FROM cards').fetchone()[0],3)
            wb=openpyxl.load_workbook(io.BytesIO(server.export_workbook(c)))
            self.assertEqual(wb['Arkusz1'].max_row,3)
            self.assertEqual(wb['Karty wycofane'].max_row,2)
        with self.assertRaises(server.Problem):
            with server.connect() as c:server.swap(c,{**self.b,'incomingVersion':1},self.user)
        with self.assertRaises(server.Problem):
            with server.connect() as c:server.register_card(c,{**self.new_card_body(),'serial':'SN1'},self.user)
    def test_calibration_confirmation_replaces_date_requirement(self):
        for cal in ('','2020-01-01'):
            with server.connect() as c:
                c.execute("UPDATE cards SET calibration=? WHERE id='a'",(cal,));c.commit()
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.swap(c,{**self.b,'calibrationVerified':False},self.user)
        with server.connect() as c:
            server.swap(c,self.b,self.user)
            self.assertEqual(c.execute("SELECT calibration FROM cards WHERE id='a'").fetchone()[0],'2020-01-01')
            self.assertIn('technik potwierdził',c.execute('SELECT note FROM events').fetchone()[0])
    def test_status_and_calibration_block_without_partial_write(self):
        for status,cal in [('nieustalony','2030-01-01'),('uszkodzona','2030-01-01')]:
            with server.connect() as c:
                c.execute('UPDATE cards SET status=?,calibration=? WHERE id=?',(status,cal,'a'));c.commit()
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.swap(c,self.b,self.user)
            with server.connect() as c:
                self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],0)
                self.assertEqual(c.execute("SELECT location FROM cards WHERE id='a'").fetchone()[0],'Magazyn')
    def test_physical_review_closure_is_audited(self):
        body=dict(id='a',version=0,status='sprawna',hall='Maintenance',location='Magazyn',reason='Sprawdzono S/N i miejsce',physicallyVerified=True)
        with self.assertRaises(server.Problem):
            with server.connect() as c:server.resolve_review(c,{**body,'physicallyVerified':False},self.user)
        with server.connect() as c:
            server.resolve_review(c,body,self.user)
            self.assertEqual(c.execute("SELECT version FROM cards WHERE id='a'").fetchone()[0],1)
            self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],1)
    def test_backup_can_be_restored(self):
        folder=self.path.parent/('backup-'+secrets.token_hex(4))
        try:
            target=server.daily_backup(self.path,folder)
            with sqlite3.connect(target,factory=server.Connection) as c:
                self.assertEqual(c.execute('PRAGMA integrity_check').fetchone()[0],'ok')
                self.assertEqual(c.execute('SELECT count(*) FROM cards').fetchone()[0],3)
                self.assertEqual(c.execute("SELECT serial FROM cards WHERE id='a'").fetchone()[0],'SN1')
        finally:
            for p in folder.iterdir():p.unlink()
            folder.rmdir()
    def test_disabled_account_sessions_revoked(self):
        with server.connect() as c:
            c.execute("INSERT INTO users(id,name,password,role) VALUES(1,'admin','x','admin'),(2,'t','x','technik')");c.commit()
            server.SESSIONS['test']={'id':2}
            server.update_user(c,dict(id=2,role='technik',active=False),{'id':1})
            self.assertNotIn('test',server.SESSIONS)
            self.assertEqual(c.execute('SELECT active FROM users WHERE id=2').fetchone()[0],0)
    def test_swap_and_replay_conflict(self):
        with server.connect() as c:
            server.swap(c,self.b,self.user)
            a=dict(c.execute("SELECT * FROM cards WHERE id='a'").fetchone());b=dict(c.execute("SELECT * FROM cards WHERE id='b'").fetchone())
            self.assertEqual((a['hall'],a['location']),('Test C4','ATE 08'));self.assertEqual(b['location'],'Kalibracja')
            with self.assertRaises(server.Problem):server.swap(c,self.b,self.user)
            c.rollback();self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],1)
    def test_undo_and_duplicate_undo(self):
        with server.connect() as c:
            event=server.swap(c,self.b,self.user);server.undo(c,{'id':event,'note':'Błędny wybór'},self.user)
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='a'").fetchone()[0],'Magazyn');self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],2)
            with self.assertRaises(server.Problem):server.undo(c,{'id':event,'note':'ponownie'},self.user)
    def test_different_models_always_blocked(self):
        with server.connect() as c:
            c.execute("UPDATE cards SET model='PXI 4071' WHERE id='a'");c.execute("UPDATE cards SET model='PXI 4070' WHERE id='b'");c.commit()
        for reason in ('','Uzasadnienie nie może omijać blokady'):
            with self.assertRaisesRegex(server.Problem,'różnych modelach'):
                with server.connect() as c:server.swap(c,{**self.b,'compatibilityReason':reason},self.user)
        with server.connect() as c:
            self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],0)
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='a'").fetchone()[0],'Magazyn')
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='b'").fetchone()[0],'ATE 08')
    def test_pxi_e_alias_matches_pxie(self):
        for value in ('PXI-E 4081','PXI-e 4081','PXIe 4081','PXIE4081',' pxi - e 4081 '):
            self.assertEqual(server.model_key(value),'PXIE4081')
        self.assertNotEqual(server.model_key('PXI 4081'),server.model_key('PXIe 4081'))
        self.assertNotEqual(server.model_key('PXI 4070'),server.model_key('PXI 4071'))
        with server.connect() as c:
            c.execute("UPDATE cards SET model='PXI-E 4081' WHERE id='a'");c.commit()
            server.swap(c,self.b,self.user)
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='a'").fetchone()[0],'ATE 08')
    def test_each_physical_confirmation_required(self):
        for key in ('incomingVerified','outgoingVerified','destinationVerified'):
            for value in (None,False,'true',1):
                body={**self.b,key:value}
                if value is None:body.pop(key)
                with self.assertRaisesRegex(server.Problem,'fizyczną weryfikację'):
                    with server.connect() as c:server.swap(c,body,self.user)
        with server.connect() as c:
            self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],0)
            server.swap(c,self.b,self.user)
            event=c.execute('SELECT * FROM events').fetchone()
            self.assertIn('Weryfikacja technika: potwierdzono',event['note'])
            self.assertEqual(event['actor'],'technik')
    def test_model_punctuation_preserves_families_and_numbers(self):
        for a,b in [('SCXI-1331','SCXI 1331'),('PXIE 2575','PXIe 2575'),('Ni 9244','NI9244'),('PXI-4071','PXI 4071')]:
            self.assertEqual(server.model_key(a),server.model_key(b))
        for a,b in [('4070','PXI 4070'),('PXI 2527','PXIe 2527'),('PXI 8364','PXIe 8364'),('NI 9375','NI 9476'),('PXIe 4081','PXIe 4083')]:
            self.assertNotEqual(server.model_key(a),server.model_key(b))
    def test_later_change_blocks_undo(self):
        with server.connect() as c:
            event=server.swap(c,self.b,self.user);c.execute("UPDATE cards SET version=version+1 WHERE id='b'");c.commit()
        with self.assertRaises(server.Problem):
            with server.connect() as c:server.undo(c,{'id':event,'note':'test'},self.user)
        with server.connect() as c:self.assertEqual(c.execute("SELECT location FROM cards WHERE id='a'").fetchone()[0],'ATE 08')
    def test_invalid_and_review_no_changes(self):
        for patch in [dict(outgoing='a'),dict(incomingVersion=99),dict(destinationLocation='ATE 08',destinationHall='Test C4'),dict(note='')]:
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.swap(c,{**self.b,**patch},self.user)
        with server.connect() as c:
            c.execute("UPDATE cards SET review='Nie odnaleziono — do weryfikacji' WHERE id='a'");c.commit()
        with self.assertRaises(server.Problem):
            with server.connect() as c:server.swap(c,self.b,self.user)
        with server.connect() as c:self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],0)
    def test_export_roundtrip_and_formula_injection(self):
        with server.connect() as c:
            server.swap(c,self.b,self.user);c.execute("UPDATE cards SET notes='=HYPERLINK(1)' WHERE id='a'");c.commit();raw=server.export_workbook(c)
        wb=openpyxl.load_workbook(io.BytesIO(raw));self.assertEqual(wb.sheetnames,['Arkusz1','Karty wycofane','Historia podmian','Informacje o eksporcie']);self.assertEqual(wb['Arkusz1'].max_row,4);self.assertEqual(wb['Historia podmian'].max_row,2)
        self.assertTrue(any(r[13].value=="'=HYPERLINK(1)" for r in wb['Arkusz1'].iter_rows(min_row=2)))
        rows=server.parse_workbook(raw);self.assertEqual(len(rows),3);self.assertEqual(next(r['location'] for r in rows if r['serial']=='SN1'),'ATE 08')
    def test_http_auth_csrf_roles(self):
        httpd=server.ThreadingHTTPServer(('127.0.0.1',0),server.Handler);t=threading.Thread(target=httpd.serve_forever,daemon=True);t.start();origin='http://127.0.0.1:'+str(httpd.server_port)
        jar=http.cookiejar.CookieJar();client=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        def req(path,b=None,csrf='',origin_header=None):
            h={'Content-Type':'application/json','X-CSRF-Token':csrf}
            if origin_header:h['Origin']=origin_header
            r=client.open(urllib.request.Request(origin+path,data=json.dumps(b).encode() if b is not None else None,headers=h));return json.load(r)
        try:
            with self.assertRaises(urllib.error.HTTPError) as e:req('/api/data')
            self.assertEqual(e.exception.code,401)
            req('/api/setup',{'name':'admin','password':'test-password-123'});csrf=req('/api/session')['csrf']
            with self.assertRaises(urllib.error.HTTPError) as e:req('/api/swap',self.b)
            self.assertEqual(e.exception.code,403)
            with self.assertRaises(urllib.error.HTTPError) as e:req('/api/swap',self.b,csrf,'https://evil.example')
            self.assertEqual(e.exception.code,403)
            req('/api/users',{'name':'viewer','password':'test-password-123','role':'odczyt'},csrf);req('/api/logout',{},csrf)
            req('/api/login',{'name':'viewer','password':'test-password-123'});csrf=req('/api/session')['csrf']
            with self.assertRaises(urllib.error.HTTPError) as e:req('/api/swap',self.b,csrf)
            self.assertEqual(e.exception.code,403);self.assertEqual(len(req('/api/data')['cards']),3)
            for path,body in [('/api/move',self.move_body()),('/api/location',dict(name='New',hall='H2')),('/api/backup',{}),('/api/station',dict(name='ATE 16',hall='H2'))]:
                with self.assertRaises(urllib.error.HTTPError) as e:req(path,body,csrf)
                self.assertEqual(e.exception.code,403)
            with self.assertRaises(urllib.error.HTTPError) as e:req('/api/backup-info')
            self.assertEqual(e.exception.code,403)
            self.assertTrue(req('/api/locations'))
        finally:httpd.shutdown();httpd.server_close();t.join()

    def test_no_mu_does_not_block_unique_serial(self):
        with server.connect() as c:
            c.execute("UPDATE cards SET mu='Brak',review='Brak MU' WHERE id='b'");c.commit()
            server.swap(c,self.b,self.user)
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='b'").fetchone()[0],'Kalibracja')
    def test_multiple_serials_identify_one_card(self):
        with server.connect() as c:
            c.execute("UPDATE cards SET serial='SN_B1 + SN_B2',mu='Brak',review='Brak MU; Wiele numerów seryjnych w wierszu' WHERE id='b'");c.commit()
            server.swap(c,self.b,self.user)
            row=c.execute("SELECT * FROM cards WHERE id='b'").fetchone()
            self.assertEqual(row['serial'],'SN_B1 + SN_B2');self.assertEqual(row['location'],'Kalibracja')
            self.assertEqual(server.effective_review(row['review']),'')
    def test_shared_serial_component_still_blocks(self):
        with server.connect() as c:
            c.execute("UPDATE cards SET serial='SN_B1 + SN_B2' WHERE id='b'");c.execute("UPDATE cards SET serial='SN_B2' WHERE id='c'");c.commit()
        with self.assertRaisesRegex(server.Problem,'innym rekordzie'):
            with server.connect() as c:server.swap(c,self.b,self.user)
    def test_discrepancy_card_with_multiple_serials(self):
        with server.connect() as c:
            c.execute("UPDATE cards SET serial='EXP1 + EXP2' WHERE id='b'");c.commit()
            result=server.discrepancy(c,dict(expected='b',expectedVersion=0,serial='FOUND1 + FOUND2',model='PXIe 4081',reason='Rzeczywista karta',expectedAbsent=True,actualVerified=True),self.user)
            self.assertEqual(c.execute('SELECT serial FROM cards WHERE id=?',(result['actual'],)).fetchone()[0],'FOUND1 + FOUND2')
    def test_discrepancy_then_swap_and_undo(self):
        with server.connect() as c:
            result=server.discrepancy(c,dict(expected='b',expectedVersion=0,actual='c',actualVersion=0,reason='Na stanowisku jest SN3 zamiast SN2',expectedAbsent=True,actualVerified=True),self.user)
            missing=dict(c.execute("SELECT * FROM cards WHERE id='b'").fetchone())
            self.assertEqual(missing['location'],'Nieustalona');self.assertIn('do weryfikacji',missing['review'])
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='c'").fetchone()[0],'ATE 08')
            event=server.swap(c,{**self.b,'outgoing':'c','outgoingVersion':1},self.user)
            server.undo(c,dict(id=event,note='Cofnięcie podmiany'),self.user)
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='c'").fetchone()[0],'ATE 08')
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='b'").fetchone()[0],'Nieustalona')
            self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],3)
    def test_discrepancy_new_card_without_mu(self):
        with server.connect() as c:
            result=server.discrepancy(c,dict(expected='b',expectedVersion=0,serial='SN_NEW',model='PXIe 4081',mu='',reason='Nieujęta wcześniej karta',expectedAbsent=True,actualVerified=True),self.user)
            row=c.execute('SELECT * FROM cards WHERE id=?',(result['actual'],)).fetchone()
            self.assertEqual(row['mu'],'');self.assertEqual(row['location'],'ATE 08')
            c.execute('UPDATE cards SET status=?,calibration=? WHERE id=?',('sprawna','2030-01-01',result['actual']));c.commit()
            server.swap(c,{**self.b,'outgoing':result['actual'],'outgoingVersion':1},self.user)
    def test_discrepancy_conflict_and_duplicate_rollback(self):
        base=dict(expected='b',expectedVersion=0,actual='c',actualVersion=0,reason='Rozbieżność',expectedAbsent=True,actualVerified=True)
        for patch in [dict(expectedVersion=9),dict(actualVersion=9),dict(expectedAbsent=False),dict(actualVerified=False),dict(actual='b')]:
            with self.assertRaises(server.Problem):
                with server.connect() as c:server.discrepancy(c,{**base,**patch},self.user)
        with self.assertRaises(server.Problem):
            with server.connect() as c:server.discrepancy(c,{**base,'actual':None,'serial':'SN1','model':'PXIe 4081'},self.user)
        with server.connect() as c:
            self.assertEqual(c.execute('SELECT count(*) FROM events').fetchone()[0],0)
            self.assertEqual(c.execute("SELECT location FROM cards WHERE id='b'").fetchone()[0],'ATE 08')

if __name__=='__main__':unittest.main(verbosity=2)
