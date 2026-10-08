"""Maintenance: rejestr kart. Python 3.11+, openpyxl. Uruchom: python server.py."""
import argparse, base64, datetime as dt, hashlib, hmac, io, json, os, re, secrets, sqlite3, threading, time
from collections import Counter
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from pathlib import Path
from http.cookies import SimpleCookie
from urllib.parse import urlparse
import openpyxl

ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / 'data' / 'maintenance.sqlite3'
SESSIONS = {}
ATTEMPTS = {}
LOCK = threading.Lock()
SECURE = False
BACKUP_LOCK = threading.Lock()
BACKUP_ERROR = ''
STATUSES = ('nieustalony','sprawna','do sprawdzenia','kalibracja','naprawa','uszkodzona','wycofana')

class Connection(sqlite3.Connection):
    def __exit__(self, *args):
        try: return super().__exit__(*args)
        finally: self.close()

def connect():
    c = sqlite3.connect(DB_PATH, timeout=15, factory=Connection)
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    return c

def init_db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with connect() as c:
        c.execute('PRAGMA journal_mode=WAL')
        c.executescript('''
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL, password TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('admin','technik','odczyt')));
        CREATE TABLE IF NOT EXISTS cards(id TEXT PRIMARY KEY, mu TEXT NOT NULL, serial TEXT NOT NULL, manufacturer TEXT NOT NULL, model TEXT NOT NULL, hall TEXT NOT NULL, location TEXT NOT NULL, designation TEXT NOT NULL, calibration TEXT NOT NULL, calibration_note TEXT NOT NULL, replacement TEXT NOT NULL, notes TEXT NOT NULL, source_row INTEGER, review TEXT NOT NULL, version INTEGER NOT NULL DEFAULT 0);
        CREATE TABLE IF NOT EXISTS events(id TEXT PRIMARY KEY, at TEXT NOT NULL, actor TEXT NOT NULL, kind TEXT NOT NULL, incoming TEXT, outgoing TEXT, before_json TEXT NOT NULL, after_json TEXT NOT NULL, note TEXT NOT NULL, reverses TEXT UNIQUE);
        CREATE TABLE IF NOT EXISTS imports(id TEXT PRIMARY KEY, at TEXT NOT NULL, actor TEXT NOT NULL, filename TEXT NOT NULL, count INTEGER NOT NULL);
        ''')
        if 'status' not in {r[1] for r in c.execute('PRAGMA table_info(cards)')}:
            c.execute("ALTER TABLE cards ADD COLUMN status TEXT NOT NULL DEFAULT 'nieustalony'")
        if 'active' not in {r[1] for r in c.execute('PRAGMA table_info(users)')}:
            c.execute('ALTER TABLE users ADD COLUMN active INTEGER NOT NULL DEFAULT 1')
        if 'retired' not in {r[1] for r in c.execute('PRAGMA table_info(cards)')}:
            c.execute('ALTER TABLE cards ADD COLUMN retired INTEGER NOT NULL DEFAULT 0')
        exists=c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='locations'").fetchone()
        c.execute('CREATE TABLE IF NOT EXISTS locations(id INTEGER PRIMARY KEY, name TEXT NOT NULL, hall TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1, UNIQUE(name,hall))')
        if not exists:
            c.execute("INSERT OR IGNORE INTO locations(name,hall) SELECT DISTINCT location,hall FROM cards WHERE location<>'' AND hall<>'' AND location<>'Nieustalona'")
            c.executemany('INSERT OR IGNORE INTO locations(name,hall) VALUES(?,?)',[('Magazyn MD','Maintenance'),('Dział kalibracji','H2')])

def daily_backup(db_path=None, folder=None):
    source=Path(db_path or DB_PATH)
    folder=Path(folder or source.parent.parent/'backups');folder.mkdir(parents=True,exist_ok=True)
    target=folder/('automatyczna-'+dt.date.today().isoformat()+'.sqlite3')
    if target.exists():return target
    temporary=target.with_suffix('.tmp')
    try:
        with sqlite3.connect(source,factory=Connection) as src, sqlite3.connect(temporary,factory=Connection) as dst:
            src.backup(dst)
            if dst.execute('PRAGMA integrity_check').fetchone()[0]!='ok':raise RuntimeError('Błąd kontroli kopii bazy')
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    return target

def backup_loop():
    global BACKUP_ERROR
    while True:
        try:
            with BACKUP_LOCK:daily_backup()
            BACKUP_ERROR=''
        except Exception as e:
            BACKUP_ERROR=str(e);print('BŁĄD KOPII BAZY:',str(e),flush=True)
        threading.Event().wait(3600)

def manual_backup():
    global BACKUP_ERROR
    folder=DB_PATH.parent.parent/'backups';folder.mkdir(parents=True,exist_ok=True)
    target=folder/('reczna-'+dt.datetime.now().strftime('%Y%m%d-%H%M%S')+'-'+secrets.token_hex(3)+'.sqlite3')
    try:
        with BACKUP_LOCK, sqlite3.connect(DB_PATH,factory=Connection) as src, sqlite3.connect(target,factory=Connection) as dst:
            src.backup(dst)
            if dst.execute('PRAGMA integrity_check').fetchone()[0]!='ok':raise RuntimeError('Błąd integralności kopii')
        BACKUP_ERROR='';return {'name':target.name,'at':dt.datetime.fromtimestamp(target.stat().st_mtime,dt.timezone.utc).isoformat()}
    except Exception as e:
        BACKUP_ERROR=str(e);target.unlink(missing_ok=True);raise Problem('Nie udało się wykonać kopii bazy.',500)

def backup_info():
    files=list((DB_PATH.parent.parent/'backups').glob('*.sqlite3'))
    latest=max(files,key=lambda p:p.stat().st_mtime) if files else None
    return {'name':latest.name if latest else '', 'at':dt.datetime.fromtimestamp(latest.stat().st_mtime,dt.timezone.utc).isoformat() if latest else '', 'error':BACKUP_ERROR}

def now():
    return dt.datetime.now(dt.timezone.utc).isoformat()

def password_hash(password, salt=None):
    salt = salt or secrets.token_hex(16)
    return salt + ':' + hashlib.pbkdf2_hmac('sha256', password.encode(), salt.encode(), 310000).hex()

def text(v):
    return '' if v is None else str(v).strip()

def parse_workbook(raw):
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    sheet = next((s for s in wb if text(s.cell(1,3).value)=='Numer MU' and text(s.cell(1,4).value)=='Numer seryjny'), None)
    if sheet is None:
        raise ValueError('Wymagany układ spisu Maintenance: C Numer MU, D Numer seryjny, E Producent, F Model, G hala, H lokalizacja.')
    rows=[]
    for n, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        vals=list(row)+[None]*18
        if not text(vals[5]) or not (text(vals[2]) or text(vals[3])):
            continue
        date=vals[9]
        cal=date.date().isoformat() if isinstance(date, dt.datetime) else date.isoformat() if isinstance(date,dt.date) else text(date)
        rows.append(dict(id=secrets.token_hex(12),mu=text(vals[2]),serial=text(vals[3]),manufacturer=text(vals[4]),model=text(vals[5]),hall=text(vals[6]),location=text(vals[7]),designation=text(vals[8]),calibration=cal,calibration_note=text(vals[11]),replacement=text(vals[12]),notes=text(vals[13]),source_row=n,review='',version=0))
    if not rows: raise ValueError('Nie znaleziono kart w arkuszu.')
    counts=Counter((r['mu'].upper().replace(' ',''),r['serial'].upper()) for r in rows)
    serials=Counter(r['serial'].upper() for r in rows if r['serial'])
    for r in rows:
        issues=[]
        if not r['serial']: issues.append('Brak numeru seryjnego')
        if serials[r['serial'].upper()]>1: issues.append('Powtórzony numer seryjny')
        if counts[(r['mu'].upper().replace(' ',''),r['serial'].upper())]>1: issues.append('Powtórzona para MU/SN')
        if not r['location']: issues.append('Brak lokalizacji')
        r['review']='; '.join(issues)
    return rows

def insert_cards(c, rows):
    for r in rows:
        keys=list(r)
        c.execute('INSERT INTO cards ('+','.join(keys)+') VALUES ('+','.join('?' for _ in keys)+')',list(r.values()))
        if r.get('hall') and r.get('location') and r['location']!='Nieustalona':
            c.execute('INSERT OR IGNORE INTO locations(name,hall) VALUES(?,?)',(r['location'],r['hall']))

class Problem(Exception):
    def __init__(self, message, status=400): self.message,self.status=message,status

def required_string(b,key,limit=200):
    v=b.get(key)
    if not isinstance(v,str) or not v.strip() or len(v)>limit: raise Problem('Uzupełnij pole: '+key)
    return v.strip()

def model_key(model):
    value=''.join(model.upper().split())
    value=re.sub(r'^PXI-E(?=\d)', 'PXIE', value)
    return re.sub(r'^(PXIE|PXI|NI|SCXI|TB|WF|CB)-(?=\d)', r'\1', value)

def effective_review(value):
    return '; '.join(x.strip() for x in value.split(';') if x.strip() and x.strip() not in ('Brak MU','Wiele numerów seryjnych w wierszu'))

def require_identity(c, card, allow_location=False):
    if dict(card).get('retired'):raise Problem('Karta jest wycofana z aktywnego spisu.')
    reasons=effective_review(card['review'])
    if allow_location:
        reasons='; '.join(x.strip() for x in reasons.split(';') if x.strip()!='Brak lokalizacji')
    if reasons: raise Problem('Karta wymaga weryfikacji: '+reasons)
    serial=''.join(card['serial'].upper().split())
    tokens={''.join(t.upper().split()) for t in card['serial'].split('+')}
    if not tokens or any(t in ('','BRAK','..','-') for t in tokens): raise Problem('Wymagane prawidłowe numery seryjne karty.')
    for other in c.execute('SELECT id,serial FROM cards WHERE id<>?',(card['id'],)):
        if tokens & {''.join(t.upper().split()) for t in other['serial'].split('+')}:
            raise Problem('Ten numer seryjny występuje w innym rekordzie. Wyjaśnij duplikat przed podmianą.')

def discrepancy(c,b,user):
    c.execute('BEGIN IMMEDIATE')
    expected=c.execute('SELECT * FROM cards WHERE id=?',(required_string(b,'expected'),)).fetchone()
    if not expected: raise Problem('Nie znaleziono oczekiwanej karty.',404)
    expected=dict(expected)
    if expected.get('retired'):raise Problem('Oczekiwana karta jest wycofana.')
    if expected['version']!=b.get('expectedVersion'):raise Problem('Dane oczekiwanej karty zmieniły się. Odśwież listę.',409)
    if not expected['location'].upper().startswith(('ATE','SATE','VATE')):raise Problem('Wybierz kartę oczekiwaną na stanowisku.')
    if b.get('expectedAbsent') is not True or b.get('actualVerified') is not True:raise Problem('Potwierdź brak oczekiwanej karty oraz S/N i lokalizację znalezionej.')
    reason=required_string(b,'reason',2000)
    actual_id=b.get('actual')
    if actual_id:
        found=c.execute('SELECT * FROM cards WHERE id=?',(actual_id,)).fetchone()
        if not found:raise Problem('Nie znaleziono wskazanej karty.',404)
        actual=dict(found)
        if actual['id']==expected['id']:raise Problem('Karta znaleziona i oczekiwana muszą być różne.')
        if actual['version']!=b.get('actualVersion'):raise Problem('Dane znalezionej karty zmieniły się. Odśwież listę.',409)
        require_identity(c,actual,allow_location=True)
    else:
        serial=required_string(b,'serial');model=required_string(b,'model')
        actual=dict(id=secrets.token_hex(12),mu=text(b.get('mu')),serial=serial,manufacturer=text(b.get('manufacturer')),model=model,hall='',location='',designation='',calibration='',calibration_note='',replacement='',notes='',source_row=None,review='',version=0)
        require_identity(c,actual,allow_location=True)
        insert_cards(c,[actual])
    missing={**expected,'hall':'','location':'Nieustalona','status':'do sprawdzenia','review':'; '.join(filter(None,[effective_review(expected['review']),'Nie odnaleziono na '+expected['hall']+' / '+expected['location']+' — do weryfikacji'])),'version':expected['version']+1}
    located={**actual,'hall':expected['hall'],'location':expected['location'],'review':effective_review(actual['review']).replace('Brak lokalizacji','').strip('; '),'version':actual['version']+1}
    for r in (missing,located):c.execute('UPDATE cards SET hall=?,location=?,review=?,status=?,version=? WHERE id=?',(r['hall'],r['location'],r['review'],r.get('status','nieustalony'),r['version'],r['id']))
    event=secrets.token_hex(12)
    c.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?,?,NULL)',(event,now(),user['name'],'rozbieżność na stanowisku',expected['id'],actual['id'],json.dumps([expected,actual]),json.dumps([missing,located]),reason+'\nTechnik potwierdził: oczekiwanej karty nie odnaleziono, sprawdzono S/N i fizyczną lokalizację znalezionej. To korekta stanu zastanego, nie podmiana.'))
    c.commit();return {'id':event,'actual':actual['id']}

def swap(c,b,user):
    c.execute('BEGIN IMMEDIATE')
    ids=[required_string(b,'incoming'),required_string(b,'outgoing')]
    if ids[0]==ids[1]: raise Problem('Wybierz dwie różne karty.')
    cards=[c.execute('SELECT * FROM cards WHERE id=?',(i,)).fetchone() for i in ids]
    if any(r is None for r in cards): raise Problem('Karta nie istnieje.',404)
    a,z=[dict(r) for r in cards]
    if a['version']!=b.get('incomingVersion') or z['version']!=b.get('outgoingVersion'): raise Problem('Ktoś zmienił te karty. Odśwież dane i ponownie sprawdź podmianę.',409)
    require_identity(c,a); require_identity(c,z)
    if a.get('status')!='sprawna':raise Problem('Do montażu wybierz kartę o potwierdzonym statusie sprawna.')
    destination_status=b.get('destinationStatus')
    if destination_status not in STATUSES:raise Problem('Wybierz status wyjmowanej karty.')
    if not z['location'].upper().startswith(('ATE','SATE','VATE')): raise Problem('Wyjmowana karta musi znajdować się na stanowisku ATE, SATE lub VATE.')
    if (a['hall'],a['location'])==(z['hall'],z['location']): raise Problem('Karty są już w tej samej lokalizacji.')
    model_a=model_key(a['model']); model_z=model_key(z['model'])
    if not model_a or model_a!=model_z: raise Problem('Nie można podmienić kart o różnych modelach: '+a['model']+' / '+z['model']+'. Wybierz kartę tego samego modelu.')
    for key,label in [('incomingVerified','numer seryjny i miejsce pobrania wkładanej karty'),('outgoingVerified','numer seryjny i obecną lokalizację wyjmowanej karty'),('destinationVerified','nowe miejsce odłożenia wyjmowanej karty')]:
        if b.get(key) is not True: raise Problem('Potwierdź fizyczną weryfikację: '+label+'.')
    hall=required_string(b,'destinationHall'); loc=required_string(b,'destinationLocation')
    if b.get('locationId') is not None:
        hall,loc=destination(c,b)
    if (hall,loc)==(z['hall'],z['location']): raise Problem('Wybierz inne miejsce docelowe wyjmowanej karty.')
    note=required_string(b,'note',2000)
    if b.get('calibrationVerified') is not True:raise Problem('Potwierdź, że karta wkładana ma ważną kalibrację lub nie wymaga kalibracji.')
    after_a={**a,'hall':z['hall'],'location':z['location'],'version':a['version']+1}
    after_z={**z,'hall':hall,'location':loc,'status':destination_status,'version':z['version']+1}
    for r in (after_a,after_z): c.execute('UPDATE cards SET hall=?,location=?,status=?,version=? WHERE id=?',(r['hall'],r['location'],r.get('status','nieustalony'),r['version'],r['id']))
    note+='\nKalibracja: technik potwierdził ważność kalibracji lub brak wymogu kalibracji karty wkładanej. Data w bazie pozostaje bez zmiany.'
    note+='\nWeryfikacja technika: potwierdzono S/N i miejsce pobrania wkładanej karty, S/N i lokalizację wyjmowanej karty oraz nowe miejsce jej odłożenia.'
    event=secrets.token_hex(12)
    c.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?,?,NULL)',(event,now(),user['name'],'podmiana',a['id'],z['id'],json.dumps([a,z],ensure_ascii=False),json.dumps([after_a,after_z],ensure_ascii=False),note))
    c.commit()
    return event

def undo(c,b,user):
    c.execute('BEGIN IMMEDIATE')
    e=c.execute('SELECT * FROM events WHERE id=? AND kind=?',(b.get('id'),'podmiana')).fetchone()
    if not e: raise Problem('Nie znaleziono podmiany.',404)
    if c.execute('SELECT 1 FROM events WHERE reverses=?',(e['id'],)).fetchone(): raise Problem('Podmiana została już cofnięta.',409)
    before=json.loads(e['before_json']); after=json.loads(e['after_json']); current=[]; restored=[]
    for old, expected in zip(before,after):
        r=c.execute('SELECT * FROM cards WHERE id=?',(old['id'],)).fetchone()
        if not r or r['version']!=expected['version']: raise Problem('Karta została później zmieniona. Cofnięcie tej podmiany jest zablokowane.',409)
        current.append(dict(r)); new={**old,'version':r['version']+1}; restored.append(new)
        c.execute('UPDATE cards SET hall=?,location=?,status=?,version=? WHERE id=?',(new['hall'],new['location'],new.get('status','nieustalony'),new['version'],new['id']))
    c.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?,?,?)',(secrets.token_hex(12),now(),user['name'],'cofnięcie',e['incoming'],e['outgoing'],json.dumps(current),json.dumps(restored),required_string(b,'note',2000),e['id']))
    c.commit()

def destination(c,b):
    row=c.execute('SELECT * FROM locations WHERE id=? AND active=1',(b.get('locationId'),)).fetchone()
    if not row:raise Problem('Wybierz aktywne miejsce ze słownika lokalizacji.')
    return row['hall'],row['name']

def location_update(c,b,user):
    c.execute('BEGIN IMMEDIATE')
    name=required_string(b,'name');hall=required_string(b,'hall')
    active=b.get('active',True)
    if type(active) is not bool:raise Problem('Nieprawidłowa aktywność miejsca.')
    if b.get('id'):
        old=c.execute('SELECT * FROM locations WHERE id=?',(b['id'],)).fetchone()
        if not old:raise Problem('Nie znaleziono miejsca.',404)
        if (name,hall)!=(old['name'],old['hall']):raise Problem('Aby zmienić nazwę, dodaj nowe miejsce. Historia zachowuje dawny zapis.')
        c.execute('UPDATE locations SET active=? WHERE id=?',(int(active),b['id']))
    else:
        if c.execute('SELECT 1 FROM locations WHERE LOWER(name)=LOWER(?) AND LOWER(hall)=LOWER(?)',(name,hall)).fetchone():raise Problem('Takie miejsce już istnieje.',409)
        c.execute('INSERT INTO locations(name,hall,active) VALUES(?,?,?)',(name,hall,int(active)))
    c.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?,?,NULL)',(secrets.token_hex(12),now(),user['name'],'słownik lokalizacji',None,None,'[]','[]',name+' / '+hall+(' — aktywne' if active else ' — wyłączone')))
    c.commit();return {'ok':True}

def move_card(c,b,user):
    c.execute('BEGIN IMMEDIATE')
    row=c.execute('SELECT * FROM cards WHERE id=?',(required_string(b,'id'),)).fetchone()
    if not row:raise Problem('Nie znaleziono karty.',404)
    old=dict(row)
    if old['version']!=b.get('version'):raise Problem('Karta zmieniła się. Odśwież dane.',409)
    require_identity(c,old)
    if b.get('physicallyVerified') is not True:raise Problem('Potwierdź S/N, obecne miejsce i miejsce docelowe.')
    hall,location=destination(c,b)
    status=b.get('status')
    if status not in STATUSES or status=='wycofana':raise Problem('Wybierz status aktywnej karty. Do wycofania użyj osobnej operacji.')
    operation=b.get('operation','przeniesienie')
    if operation not in ('przeniesienie','powrót z kalibracji','powrót z naprawy'):raise Problem('Wybierz rodzaj operacji.')
    reason=required_string(b,'reason',2000)
    calibration=old['calibration'];cal_note=old['calibration_note']
    if operation!='przeniesienie':
        result=b.get('result');document=required_string(b,'document',500)
        if result not in ('pozytywny','negatywny'):raise Problem('Wybierz wynik usługi.')
        if result=='negatywny' and status=='sprawna':raise Problem('Po negatywnym wyniku karta nie może otrzymać statusu sprawna.')
        calibration=b.get('calibration',old['calibration'])
        if not isinstance(calibration,str) or len(calibration)>100:raise Problem('Nieprawidłowa data kalibracji.')
        calibration=calibration.strip()
        if calibration and calibration.lower()!='nie dotyczy':
            try:dt.date.fromisoformat(calibration)
            except ValueError:raise Problem('Kalibracja: RRRR-MM-DD lub Nie dotyczy.')
        cal_note=operation+'; wynik: '+result+'; dokument: '+document
        reason+='\n'+cal_note
    if location.upper().startswith(('ATE','SATE','VATE')):
        if status!='sprawna' or b.get('calibrationVerified') is not True:raise Problem('Montaż na ATE wymaga sprawności i potwierdzenia kalibracji.')
        if b.get('installationVerified') is not True:raise Problem('Potwierdź montaż bez zastępowania innej karty; dla wymiany użyj podmiany.')
    if (hall,location,status)==(old['hall'],old['location'],old['status']) and operation=='przeniesienie':raise Problem('Wybierz nowe miejsce lub zmieniony status.')
    new={**old,'hall':hall,'location':location,'status':status,'calibration':calibration,'calibration_note':cal_note,'version':old['version']+1}
    c.execute('UPDATE cards SET hall=?,location=?,status=?,calibration=?,calibration_note=?,version=? WHERE id=?',(hall,location,status,calibration,cal_note,new['version'],old['id']))
    c.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?,?,NULL)',(secrets.token_hex(12),now(),user['name'],operation,old['id'],None,json.dumps([old]),json.dumps([new]),reason+'\nPotwierdzono fizycznie S/N i oba miejsca.'+(' Potwierdzono kalibrację i nowy montaż na stanowisku.' if location.upper().startswith(('ATE','SATE','VATE')) else '')))
    c.commit();return {'ok':True}

def register_card(c,b,user):
    c.execute('BEGIN IMMEDIATE')
    card={key:'' for key in ('mu','serial','manufacturer','model','hall','location','designation','calibration','calibration_note','replacement','notes','review')}
    for key in card:
        if key in ('review',):continue
        value=b.get(key,'')
        if not isinstance(value,str) or len(value)>2000:raise Problem('Nieprawidłowe pole: '+key)
        card[key]=value.strip()
    for key in ('serial','model','hall','location'):card[key]=required_string(b,key,2000)
    if b.get('locationId') is not None:card['hall'],card['location']=destination(c,b)
    status=b.get('status','nieustalony')
    if status not in STATUSES or status=='wycofana':raise Problem('Wybierz status nowej karty.')
    card.update(id=secrets.token_hex(12),source_row=None,version=0,status=status,retired=0)
    require_identity(c,card)
    if b.get('physicallyVerified') is not True:raise Problem('Potwierdź S/N i fizyczną lokalizację nowej karty.')
    reason=required_string(b,'reason',2000)
    cal=card['calibration']
    if cal and cal.lower()!='nie dotyczy':
        try:dt.date.fromisoformat(cal)
        except ValueError:raise Problem('Kalibracja: wpisz RRRR-MM-DD lub Nie dotyczy.')
    if card['location'].upper().startswith(('ATE','SATE','VATE')):
        if status!='sprawna':raise Problem('Nowa karta montowana na ATE musi mieć status sprawna.')
        if cal.lower()!='nie dotyczy' and (not cal or dt.date.fromisoformat(cal)<dt.date.today()):raise Problem('Montaż na ATE wymaga aktualnej kalibracji lub Nie dotyczy.')
        if b.get('installationVerified') is not True:raise Problem('Potwierdź nowy montaż na ATE. Przy wymianie istniejącej karty użyj podmiany.')
    insert_cards(c,[card])
    c.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?,?,NULL)',(secrets.token_hex(12),now(),user['name'],'dodanie karty',card['id'],None,'[]',json.dumps([card]),reason))
    c.commit();return {'id':card['id']}

def retire_card(c,b,user):
    c.execute('BEGIN IMMEDIATE')
    old=c.execute('SELECT * FROM cards WHERE id=?',(required_string(b,'id'),)).fetchone()
    if not old:raise Problem('Nie znaleziono karty.',404)
    old=dict(old)
    if old['retired']:raise Problem('Karta już została wycofana.',409)
    if old['version']!=b.get('version'):raise Problem('Dane karty zmieniły się. Odśwież spis.',409)
    reason=required_string(b,'reason',2000)
    hall=required_string(b,'hall');location=required_string(b,'location')
    if b.get('locationId') is not None:hall,location=destination(c,b)
    if location.upper().startswith(('ATE','SATE','VATE')):raise Problem('Podaj miejsce odłożenia wycofanej karty poza stanowiskiem.')
    if b.get('physicallyVerified') is not True:raise Problem('Potwierdź wycofaną kartę i miejsce jej odłożenia.')
    new={**old,'retired':1,'status':'wycofana','hall':hall,'location':location,'version':old['version']+1}
    c.execute('UPDATE cards SET retired=1,status=?,hall=?,location=?,version=? WHERE id=?',('wycofana',hall,location,new['version'],old['id']))
    c.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?,?,NULL)',(secrets.token_hex(12),now(),user['name'],'wycofanie karty',old['id'],None,json.dumps([old]),json.dumps([new]),reason))
    c.commit();return {'ok':True}

def resolve_review(c,b,user):
    c.execute('BEGIN IMMEDIATE')
    old=c.execute('SELECT * FROM cards WHERE id=?',(required_string(b,'id'),)).fetchone()
    if not old:raise Problem('Nie znaleziono karty.',404)
    old=dict(old)
    if old['version']!=b.get('version'):raise Problem('Karta zmieniła się. Odśwież dane.',409)
    if b.get('physicallyVerified') is not True:raise Problem('Potwierdź fizyczną identyfikację karty i lokalizację.')
    require_identity(c,{**old,'review':''},allow_location=True)
    status=b.get('status')
    if status not in STATUSES:raise Problem('Wybierz status techniczny.')
    hall=required_string(b,'hall');location=required_string(b,'location')
    if b.get('locationId') is not None:hall,location=destination(c,b)
    if location.lower()=='nieustalona':raise Problem('Podaj faktycznie potwierdzoną lokalizację.')
    note=required_string(b,'reason',2000)
    calibration=b.get('calibration',old['calibration'])
    if not isinstance(calibration,str):raise Problem('Nieprawidłowa kalibracja.')
    calibration=calibration.strip()
    if calibration and calibration.lower()!='nie dotyczy':
        try:dt.date.fromisoformat(calibration)
        except ValueError:raise Problem('Kalibracja: wpisz RRRR-MM-DD lub Nie dotyczy.')
    new={**old,'hall':hall,'location':location,'review':'','status':status,'calibration':calibration,'version':old['version']+1}
    c.execute('UPDATE cards SET hall=?,location=?,review=?,status=?,calibration=?,version=? WHERE id=?',(hall,location,'',status,calibration,new['version'],old['id']))
    c.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?,?,NULL)',(secrets.token_hex(12),now(),user['name'],'zamknięcie weryfikacji',old['id'],None,json.dumps([old]),json.dumps([new]),note+'; Administrator potwierdził fizyczny S/N i miejsce.'))
    c.commit();return {'ok':True}

def update_user(c,b,user):
    c.execute('BEGIN IMMEDIATE')
    target=c.execute('SELECT * FROM users WHERE id=?',(b.get('id'),)).fetchone()
    if not target:raise Problem('Nie znaleziono konta.',404)
    role=b.get('role');active=b.get('active')
    if role not in ('admin','technik','odczyt') or type(active) is not bool:raise Problem('Nieprawidłowa rola lub aktywność.')
    if target['id']==user['id'] and (not active or role!='admin'):raise Problem('Nie możesz odebrać własnego dostępu administratora.')
    if target['role']=='admin' and target['active'] and (not active or role!='admin') and c.execute("SELECT COUNT(*) FROM users WHERE active=1 AND role='admin'").fetchone()[0]<=1:raise Problem('Musi pozostać aktywny administrator.')
    password=b.get('password','')
    if password and (not isinstance(password,str) or len(password)<12):raise Problem('Nowe hasło: minimum 12 znaków.')
    c.execute('UPDATE users SET role=?,active=? WHERE id=?',(role,int(active),target['id']))
    if password:c.execute('UPDATE users SET password=? WHERE id=?',(password_hash(password),target['id']))
    c.commit()
    with LOCK:
        for token,session in list(SESSIONS.items()):
            if session['id']==target['id']:del SESSIONS[token]
    return {'ok':True}

def export_workbook(c):
    if not c.in_transaction:c.execute('BEGIN')
    wb=openpyxl.Workbook(); s=wb.active; s.title='Arkusz1'
    s.append(['Pole techniczne','Lp.','Numer MU','Numer seryjny','Producent','Model','Lokalizacja - hala','Lokalizacja','Domyślna/zastępcza','Data końca kalibracji','Dni','Oddana do kalibracji +data','Jaka karta zastępuje tę','Uwagi 1','ID aplikacji','Weryfikacja','Status techniczny'])
    archive=wb.create_sheet('Karty wycofane');archive.append([cell.value for cell in s[1]])
    def safe(v): return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v
    for i,r in enumerate(c.execute('SELECT * FROM cards ORDER BY location,model,mu'),1):
        date=r['calibration']; days=''
        try: date=dt.date.fromisoformat(date); days=(date-dt.date.today()).days
        except ValueError: pass
        target=archive if r['retired'] else s
        target.append([safe(v) for v in ['',target.max_row,r['mu'],r['serial'],r['manufacturer'],r['model'],r['hall'],r['location'],r['designation'],date,days,r['calibration_note'],r['replacement'],r['notes'],r['id'],effective_review(r['review']),r['status']]])
        target.cell(target.max_row,10).number_format='dd.mm.yyyy'
    h=wb.create_sheet('Historia podmian'); h.append(['ID','Data UTC','Technik','Operacja','MU wkładanej','SN wkładanej','MU wyjmowanej','SN wyjmowanej','Przed wkładana','Po wkładana','Przed wyjmowana','Po wyjmowana','Uwagi','Cofa operację','Status wkładanej przed','Status wkładanej po','Status wyjmowanej przed','Status wyjmowanej po','Kalibracja wkładanej przed','Kalibracja wkładanej po','Kalibracja wyjmowanej przed','Kalibracja wyjmowanej po'])
    for e in c.execute('SELECT * FROM events ORDER BY at'):
        before=json.loads(e['before_json']); after=json.loads(e['after_json'])
        a=before[0] if before else after[0] if after else {}; z=before[1] if len(before)>1 else {}
        place=lambda r:r.get('hall','')+' / '+r.get('location','')
        h.append([safe(v) for v in [e['id'],e['at'],e['actor'],e['kind'],a.get('mu',''),a.get('serial',''),z.get('mu',''),z.get('serial',''),place(a),place(after[0]) if after else '',place(z),place(after[1]) if len(after)>1 else '',e['note'],e['reverses'] or '',a.get('status',''),after[0].get('status','') if after else '',z.get('status',''),after[1].get('status','') if len(after)>1 else '',a.get('calibration',''),after[0].get('calibration','') if after else '',z.get('calibration',''),after[1].get('calibration','') if len(after)>1 else '']])
    meta=wb.create_sheet('Informacje o eksporcie');meta.append(['Parametr','Wartość']);meta.append(['Wygenerowano UTC',now()]);meta.append(['Wersja stanu',c.execute('SELECT COUNT(*) FROM events').fetchone()[0]]);meta.append(['Liczba kart',c.execute('SELECT COUNT(*) FROM cards').fetchone()[0]]);meta.append(['Źródło','Migawka wspólnej bazy; zmiany w pliku nie aktualizują aplikacji'])
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.worksheet.table import Table,TableStyleInfo
    for sh in wb:
        sh.freeze_panes='C2' if sh==s else 'B2'; sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='123C36'); cell.alignment=Alignment(wrap_text=True)
        sh.row_dimensions[1].height=36
        for col in sh.columns:
            letter=col[0].column_letter
            sh.column_dimensions[letter].width=min(48,max(14,max(len(str(v.value or '')) for v in col)+2))
        if sh.max_row>1:
            table=Table(displayName='Karty' if sh==s else 'Wycofane' if sh==archive else 'Historia' if sh==h else 'Informacje',ref=sh.dimensions);table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);sh.add_table(table)
    out=io.BytesIO();wb.save(out);return out.getvalue()

class Handler(BaseHTTPRequestHandler):
    def log_message(self,fmt,*args): pass
    def reply(self,data,status=200,ctype='application/json; charset=utf-8',headers=None):
        raw=json.dumps(data,ensure_ascii=False).encode() if ctype.startswith('application/json') else data
        self.send_response(status);self.send_header('Content-Type',ctype);self.send_header('Content-Length',str(len(raw)));self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff');self.send_header('X-Frame-Options','DENY');self.send_header('Content-Security-Policy',"default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; frame-ancestors 'none'")
        for k,v in (headers or {}).items(): self.send_header(k,v)
        self.end_headers();self.wfile.write(raw)
    def session(self):
        try: ck=SimpleCookie(self.headers.get('Cookie','')); token=ck['session'].value
        except (KeyError,ValueError): return None
        with LOCK:
            s=SESSIONS.get(token)
            if s and s['expires']>time.time():
                with connect() as c:u=c.execute('SELECT active,role FROM users WHERE id=?',(s['id'],)).fetchone()
                if u and u['active']:
                    s['role']=u['role'];return s
        return None
    def body(self):
        size=int(self.headers.get('Content-Length','0'))
        if size>8_000_000: raise Problem('Plik jest zbyt duży (maks. 8 MB).',413)
        if self.headers.get('Content-Type','').split(';')[0]!='application/json': raise Problem('Wymagany JSON.',415)
        try: return json.loads(self.rfile.read(size))
        except Exception: raise Problem('Nieprawidłowe dane JSON.')
    def do_GET(self):
        try:
            path=urlparse(self.path).path
            if path in ('/','/app.js','/enhancements.js','/style.css'):
                filename={'/':'index.html','/app.js':'app.js','/enhancements.js':'enhancements.js','/style.css':'style.css'}[path]
                return self.reply((ROOT/'static'/filename).read_bytes(),ctype={'/':'text/html; charset=utf-8','/app.js':'text/javascript; charset=utf-8','/enhancements.js':'text/javascript; charset=utf-8','/style.css':'text/css; charset=utf-8'}[path])
            user=self.session()
            with connect() as c:
                if path=='/api/session': return self.reply({'user':{k:user[k] for k in ('id','name','role')} if user else None,'csrf':user['csrf'] if user else '', 'setup':c.execute('SELECT COUNT(*) FROM users').fetchone()[0]==0})
                if not user: raise Problem('Zaloguj się.',401)
                if path=='/api/data': return self.reply({'cards':[{**dict(r),'review':effective_review(r['review'])} for r in c.execute('SELECT * FROM cards ORDER BY location,model,mu')],'events':[dict(r) for r in c.execute('SELECT * FROM events ORDER BY at DESC')],'imports':[dict(r) for r in c.execute('SELECT * FROM imports ORDER BY at DESC')]})
                if path=='/api/users':
                    if user['role']!='admin': raise Problem('Wymagane uprawnienia administratora.',403)
                    return self.reply([dict(r) for r in c.execute('SELECT id,name,role,active FROM users')])
                if path=='/api/locations':return self.reply([dict(r) for r in c.execute('SELECT * FROM locations ORDER BY name,hall')])
                if path=='/api/backup-info':
                    if user['role']!='admin':raise Problem('Wymagany administrator.',403)
                    return self.reply(backup_info())
                if path=='/api/export': return self.reply(export_workbook(c),ctype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="Spis_kart_Maintenance.xlsx"'})
            raise Problem('Nie znaleziono.',404)
        except Problem as e:self.reply({'error':e.message},e.status)
        except Exception as e: print(type(e).__name__,str(e));self.reply({'error':'Błąd serwera. Sprawdź dziennik aplikacji.'},500)
    def do_POST(self):
        try:
            origin=self.headers.get('Origin')
            if origin and urlparse(origin).netloc!=self.headers.get('Host'): raise Problem('Niedozwolone źródło żądania.',403)
            b=self.body(); path=urlparse(self.path).path
            with connect() as c:
                if path in ('/api/setup','/api/login'):
                    name=required_string(b,'name',80);password=required_string(b,'password',256)
                    if path=='/api/setup':
                        if self.client_address[0] not in ('127.0.0.1','::1'): raise Problem('Pierwsze konto utwórz na serwerze przez localhost.',403)
                        c.execute('BEGIN IMMEDIATE')
                        if c.execute('SELECT COUNT(*) FROM users').fetchone()[0]: raise Problem('Administrator już istnieje.',409)
                        if len(password)<12: raise Problem('Hasło musi mieć co najmniej 12 znaków.')
                        c.execute('INSERT INTO users(name,password,role) VALUES(?,?,?)',(name,password_hash(password),'admin'));c.commit()
                    key=(self.client_address[0],name.lower())
                    with LOCK:
                        recent=[v for v in ATTEMPTS.get(key,[]) if v>time.time()-900]
                        if len(recent)>=10: raise Problem('Za dużo prób logowania. Spróbuj za 15 minut.',429)
                        ATTEMPTS[key]=recent+[time.time()]
                    u=c.execute('SELECT * FROM users WHERE name=?',(name,)).fetchone()
                    if not u or not u['active'] or not hmac.compare_digest(password_hash(password,u['password'].split(':')[0]),u['password']): raise Problem('Nieprawidłowy login lub hasło.',401)
                    token=secrets.token_urlsafe(32)
                    with LOCK:
                        SESSIONS[token]={**dict(u),'csrf':secrets.token_urlsafe(24),'expires':time.time()+28800};ATTEMPTS.pop(key,None)
                    return self.reply({'ok':True},headers={'Set-Cookie':'session='+token+'; HttpOnly; SameSite=Strict; Path=/; Max-Age=28800'+('; Secure' if SECURE else '')})
                user=self.session()
                if not user: raise Problem('Sesja wygasła. Zaloguj się ponownie.',401)
                if self.headers.get('X-CSRF-Token')!=user['csrf']: raise Problem('Odśwież stronę i spróbuj ponownie.',403)
                if path=='/api/logout':
                    with LOCK:
                        for k,v in list(SESSIONS.items()):
                            if v is user: del SESSIONS[k]
                    return self.reply({'ok':True},headers={'Set-Cookie':'session=; HttpOnly; SameSite=Strict; Path=/; Max-Age=0'})
                if user['role']=='odczyt': raise Problem('To konto ma tylko dostęp do odczytu.',403)
                if path=='/api/discrepancy': return self.reply(discrepancy(c,b,user))
                if path=='/api/swap': return self.reply({'id':swap(c,b,user)})
                if path=='/api/move':return self.reply(move_card(c,b,user))
                if path=='/api/undo':
                    if user['role']!='admin': raise Problem('Tylko administrator może cofnąć podmianę.',403)
                    undo(c,b,user);return self.reply({'ok':True})
                if user['role']!='admin': raise Problem('Wymagane uprawnienia administratora.',403)
                if path=='/api/location':return self.reply(location_update(c,b,user))
                if path=='/api/backup':return self.reply(manual_backup())
                if path=='/api/resolve':return self.reply(resolve_review(c,b,user))
                if path=='/api/card-add':return self.reply(register_card(c,b,user))
                if path=='/api/card-retire':return self.reply(retire_card(c,b,user))
                if path=='/api/user-update':return self.reply(update_user(c,b,user))
                if path=='/api/users':
                    name=required_string(b,'name',80);password=required_string(b,'password',256);role=b.get('role')
                    if role not in ('admin','technik','odczyt') or len(password)<12: raise Problem('Wybierz rolę i hasło z co najmniej 12 znakami.')
                    try:c.execute('INSERT INTO users(name,password,role) VALUES(?,?,?)',(name,password_hash(password),role));c.commit()
                    except sqlite3.IntegrityError:raise Problem('Taki login już istnieje.',409)
                    return self.reply({'ok':True})
                if path in ('/api/import-preview','/api/import'):
                    try: raw=base64.b64decode(b.get('file',''),validate=True);rows=parse_workbook(raw)
                    except Exception as e: raise Problem('Nie można odczytać pliku: '+str(e))
                    if path.endswith('preview'):return self.reply({'count':len(rows),'reviewCount':sum(bool(r['review']) for r in rows),'sample':rows[:8]})
                    c.execute('BEGIN IMMEDIATE')
                    if c.execute('SELECT COUNT(*) FROM cards').fetchone()[0] or c.execute('SELECT COUNT(*) FROM events').fetchone()[0]:raise Problem('Baza zawiera dane. Import początkowy nie może nadpisać kart ani historii.',409)
                    insert_cards(c,rows);c.execute('INSERT INTO imports VALUES(?,?,?,?,?)',(secrets.token_hex(12),now(),user['name'],text(b.get('filename')),len(rows)));c.commit();return self.reply({'count':len(rows)})
                if path=='/api/card':
                    note=required_string(b,'reason',2000);c.execute('BEGIN IMMEDIATE');old=c.execute('SELECT * FROM cards WHERE id=?',(b.get('id'),)).fetchone()
                    if not old:raise Problem('Nie znaleziono karty.',404)
                    old=dict(old)
                    if old['retired']:raise Problem('Wycofana karta pozostaje w archiwum. Nie można jej edytować.')
                    if old['version']!=b.get('version'):raise Problem('Karta została zmieniona. Odśwież dane.',409)
                    fields=['mu','serial','manufacturer','model','designation','calibration','calibration_note','replacement','notes','status']
                    new={**old}
                    for f in fields:
                        v=b.get(f)
                        if not isinstance(v,str) or len(v)>2000: raise Problem('Nieprawidłowe pole '+f)
                        new[f]=v.strip()
                    if new['status'] not in STATUSES:raise Problem('Nieprawidłowy status techniczny.')
                    if not new['serial'] or not new['model']:raise Problem('Numer seryjny i model są wymagane.')
                    if new['calibration'] not in ('','Nie dotyczy'):
                        try:dt.date.fromisoformat(new['calibration'])
                        except ValueError:raise Problem('Kalibracja: wpisz RRRR-MM-DD lub Nie dotyczy.')
                    if b.get('resolveReview') is True:
                        if c.execute('SELECT 1 FROM cards WHERE UPPER(serial)=UPPER(?) AND id<>?',(new['serial'],old['id'])).fetchone(): raise Problem('Numer seryjny nadal występuje w innym rekordzie.')
                        require_identity(c,{**new,'review':''},allow_location=True)
                        new['review']=''
                    new['version']+=1
                    changed=fields+['review','version'];c.execute('UPDATE cards SET '+','.join(f+'=?' for f in changed)+' WHERE id=?',[new[f] for f in changed]+[old['id']])
                    c.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?,?,NULL)',(secrets.token_hex(12),now(),user['name'],'korekta',old['id'],None,json.dumps([old]),json.dumps([new]),note));c.commit();return self.reply({'ok':True})
                raise Problem('Nie znaleziono.',404)
        except Problem as e:self.reply({'error':e.message},e.status)
        except Exception as e:print(type(e).__name__,str(e));self.reply({'error':'Błąd serwera. Sprawdź dziennik aplikacji.'},500)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--host',default='127.0.0.1');p.add_argument('--port',type=int,default=8765);p.add_argument('--db');p.add_argument('--secure-cookie',action='store_true');a=p.parse_args()
    if a.db:DB_PATH=Path(a.db).resolve()
    SECURE=a.secure_cookie;init_db()
    daily_backup()
    threading.Thread(target=backup_loop,daemon=True).start()
    print(f'Maintenance: http://{a.host}:{a.port} (Ctrl+C zamyka serwer)',flush=True)
    ThreadingHTTPServer((a.host,a.port),Handler).serve_forever()
